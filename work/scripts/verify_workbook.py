"""Phase 6 — headless LibreOffice recalc, then verify: zero formula errors; every CHECK row = 0 (or n/a)."""
import subprocess, sys, os, shutil, glob
from openpyxl import load_workbook
W = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = f'{W}/deliverables/MBGL_Historical_Model.xlsx'
tmp = '/tmp/claude-0/-home-user-MBGL/ea5d3d5d-c379-5607-9c68-2fb9c6b18061/scratchpad/recalc'
shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, src], check=True, capture_output=True, timeout=180)
wb = load_workbook(f'{tmp}/MBGL_Historical_Model.xlsx', data_only=True)
errs = []; checks = 0; fails = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and c.value[1:4] in ('DIV', 'REF', 'VAL', 'NAM', 'N/A', 'NUM', 'NUL'):
                errs.append(f'{ws.title}!{c.coordinate} {c.value}')
        if ws.cell(row[0].row, 26).value == 'CHECK':
            for c in row[1:25]:
                if isinstance(c.value, (int, float)):
                    checks += 1
                    if abs(c.value) > 0.5: fails.append(f'{ws.title}!{c.coordinate} {ws.cell(c.row,1).value} = {c.value}')
print('formula errors:', len(errs)); [print(' ', e) for e in errs[:30]]
print('check cells evaluated:', checks, '| failures:', len(fails)); [print(' ', f) for f in fails]
sys.exit(1 if errs or fails else 0)
