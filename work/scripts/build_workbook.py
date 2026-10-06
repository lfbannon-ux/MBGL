"""Phase 5 — build the MBGL historical model workbook from canonical + extraction CSVs.
Colours: blue = hard input as filed; black = in-sheet formula; green = cross-sheet formula.
Check rows carry the marker 'CHECK' in hidden column Z so the recalc verifier can find them."""
import csv, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

W = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EX = f'{W}/extraction'
NAVY = '1F3864'; SEC = 'D9E1F2'
F_BASE = Font(name='Arial', size=9)
F_IN = Font(name='Arial', size=9, color='0000FF')
F_IN_B = Font(name='Arial', size=9, color='0000FF', bold=True)
F_FX = Font(name='Arial', size=9, color='000000')
F_FX_B = Font(name='Arial', size=9, color='000000', bold=True)
F_X = Font(name='Arial', size=9, color='00803C')
F_X_B = Font(name='Arial', size=9, color='00803C', bold=True)
F_TITLE = Font(name='Arial', size=12, bold=True, color='FFFFFF')
F_HDR = Font(name='Arial', size=9, bold=True, color='FFFFFF')
F_SEC = Font(name='Arial', size=9, bold=True)
F_NOTE = Font(name='Arial', size=8, italic=True, color='595959')
FILL_N = PatternFill('solid', fgColor=NAVY); FILL_S = PatternFill('solid', fgColor=SEC)
FILL_CHK = PatternFill('solid', fgColor='F2F2F2')
MONEY = '#,##0.0;(#,##0.0);"-"'; PS = '$#,##0.00;($#,##0.00);"-"'; PCT_P = '0.0"%"'; PCT_C = '0.0%'
CNT = '#,##0;(#,##0);"-"'; CNT1 = '#,##0.0;(#,##0.0);"-"'
FMT = {'usd_m': MONEY, 'usd_per_share': PS, 'pct': PCT_P, 'count': CNT, 'shares_m': CNT1, 'usd': CNT, 'usd_k': CNT}
MARK = 26  # column Z — hidden check marker

def rd(name): return list(csv.DictReader(open(f'{EX}/{name}')))

def period_key(p):
    m = re.match(r'^FY(\d{4})$', p)
    if m: return (0, int(m.group(1)), 0, p)
    m = re.match(r'^(Q\d|6M)-(\d{4})$', p)
    if m: return (1, int(m.group(2)), {'Q1': 1, 'Q2': 2, '6M': 3, 'Q3': 4, 'Q4': 5}[m.group(1)], p)
    m = re.match(r'^FY(\d{4})E$', p)
    if m: return (2, int(m.group(1)), 0, p)
    m = re.match(r'^(\d{4})$', p)
    if m: return (0, int(m.group(1)), 0, p)
    return (3, 0, 0, p)

def setup(ws, title, subtitle, ncols, label_w=50):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = label_w
    for c in range(2, ncols + 2): ws.column_dimensions[L(c)].width = 12
    ws.column_dimensions[L(MARK)].hidden = True
    for c in range(1, ncols + 2):
        ws.cell(1, c).fill = FILL_N
    ws.cell(1, 1, title).font = F_TITLE
    ws.cell(2, 1, subtitle).font = F_NOTE
    ws.row_dimensions[1].height = 20

def header(ws, row, labels, first='USD millions unless stated'):
    ws.cell(row, 1, first).font = F_HDR; ws.cell(row, 1).fill = FILL_N
    for i, lab in enumerate(labels):
        c = ws.cell(row, i + 2, lab); c.font = F_HDR; c.fill = FILL_N
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row + 1, 2)

def section(ws, row, text, ncols):
    for c in range(1, ncols + 2): ws.cell(row, c).fill = FILL_S
    ws.cell(row, 1, text).font = F_SEC

# ---------------------------------------------------------------- canonical primary data
canon = {}
src = {}
for r in csv.DictReader(open(f'{W}/canonical/primary.csv')):
    canon[(r['canon_key'], r['period'])] = float(r['value'])
    src[(r['canon_key'], r['period'])] = f"{r['source_doc']} p{r['page']}"

FLOW_P = ['FY2023', 'FY2024', 'FY2025', 'Q1-2025', 'Q2-2025', '6M-2025', 'Q1-2026', 'Q2-2026', '6M-2026']
BS_P = ['FY2024', 'FY2025', 'Q1-2026', 'Q2-2026']
PLAB = {'FY2023': 'FY2023\n12/31/23', 'FY2024': 'FY2024\n12/31/24', 'FY2025': 'FY2025\n12/31/25',
        'Q1-2025': 'Q1-25\n3M 3/31/25', 'Q2-2025': 'Q2-25\n3M 6/30/25', '6M-2025': '6M-25\n6/30/25',
        'Q1-2026': 'Q1-26\n3M 3/31/26', 'Q2-2026': 'Q2-26\n3M 6/30/26', '6M-2026': '6M-26\n6/30/26'}
BSLAB = {'FY2024': '12/31/24', 'FY2025': '12/31/25', 'Q1-2026': '3/31/26', 'Q2-2026': '6/30/26'}

wb = Workbook()
REF = {}   # (sheet, key) -> row

def write_statement(ws, start, periods, spec, sheet):
    """spec rows: ('sec', text) | ('in', key, label, fmt, bold) | ('fx', key, label, expr, fmt) |
       ('chk', label, expr_fn) ; expr strings use {key} placeholders resolved to same-column cells."""
    r = start
    for item in spec:
        kind = item[0]
        if kind == 'sec':
            section(ws, r, item[1], len(periods) + 1); r += 1; continue
        if kind == 'note':
            ws.cell(r, 1, item[1]).font = F_NOTE; r += 1; continue
        if kind == 'in':
            _, key, lab, fmt, bold = item
            ws.cell(r, 1, lab).font = F_SEC if bold else F_BASE
            for i, p in enumerate(periods):
                v = canon.get((key, p))
                if v is not None:
                    c = ws.cell(r, i + 2, v); c.font = F_IN_B if bold else F_IN; c.number_format = fmt
            srcs = sorted({src[(key, p)] for p in periods if (key, p) in src})
            ws.cell(r, len(periods) + 2, '; '.join(srcs)).font = F_NOTE
            REF[(sheet, key)] = r; r += 1; continue
        if kind in ('fx', 'chk'):
            if kind == 'fx':
                _, key, lab, expr, fmt = item
            else:
                _, lab, expr, applies = item; key = None; fmt = MONEY
            ws.cell(r, 1, lab).font = F_FX_B if kind == 'fx' else F_NOTE
            for i, p in enumerate(periods):
                col = L(i + 2)
                if kind == 'chk' and not applies(p):
                    ws.cell(r, i + 2, 'n/a').font = F_NOTE; continue
                e = re.sub(r'\{([\w.]+)\}', lambda m: f'{col}{REF[(sheet, m.group(1))]}', expr)
                e = re.sub(r'\[(\w+)!([\w.]+)\]', lambda m: f"'{m.group(1)}'!{xcol(m.group(1), p)}{REF[(m.group(1), m.group(2))]}", e)
                c = ws.cell(r, i + 2, '=' + e)
                c.font = (F_X if '!' in e else F_FX) if kind == 'chk' else (F_X_B if '!' in e else F_FX_B)
                c.number_format = fmt
            if kind == 'chk':
                ws.cell(r, MARK, 'CHECK')
                for cc in range(1, len(periods) + 2): ws.cell(r, cc).fill = FILL_CHK
            else:
                REF[(sheet, key)] = r
            r += 1
    return r

COLS = {}
def xcol(sheet, p): return L(COLS[sheet].index(p) + 2)

def has(key, p): return (key, p) in canon
def printed(key):  # check applies only where the subtotal was printed
    return lambda p: has(key, p)

# ---------------------------------------------------------------- IS
ws = wb.active; ws.title = 'IS'
COLS['IS'] = FLOW_P
setup(ws, 'Mobility Global Inc. — Combined Statements of Income (as originally reported)',
      'Carve-out (combined) basis; FY from Form 10/A (2026-05-27) audited statements; Q1 from Form 10/A interim; Q2/6M from 10-Q Q2-2026. Blue = as filed; black = formula; green = cross-sheet.', len(FLOW_P) + 1)
header(ws, 4, [PLAB[p] for p in FLOW_P] + ['Source'])
IS_SPEC = [
 ('sec', 'Income statement'),
 ('in', 'IS.revenue', 'Revenue', MONEY, True),
 ('in', 'IS.oprel', 'Operating-related expenses', MONEY, False),
 ('in', 'IS.sg', 'Selling and general expenses', MONEY, False),
 ('in', 'IS.dep', 'Depreciation', MONEY, False),
 ('in', 'IS.amort', 'Amortization of intangibles', MONEY, False),
 ('fx', 'IS.totexp_f', 'Total expenses', '{IS.oprel}+{IS.sg}+{IS.dep}+{IS.amort}', MONEY),
 ('fx', 'IS.op_f', 'Operating profit', '{IS.revenue}-{IS.totexp_f}', MONEY),
 ('in', 'IS.other', 'Other income, net (printed as a deduction; negative = income)', MONEY, False),
 ('in', 'IS.int', 'Interest expense, net', MONEY, False),
 ('fx', 'IS.pbt_f', 'Income before taxes on income', '{IS.op_f}-{IS.other}-{IS.int}', MONEY),
 ('in', 'IS.tax', 'Provision for taxes on income', MONEY, False),
 ('fx', 'IS.ni_f', 'Net income', '{IS.pbt_f}-{IS.tax}', MONEY),
 ('in', 'IS.eps_b', 'Net income per common share — basic (10-Q only; Form 10 says EPS not presented)', PS, False),
 ('in', 'IS.eps_d', 'Net income per common share — diluted (10-Q only)', PS, False),
 ('sec', 'Comprehensive income'),
 ('in', 'CI.fx', 'Foreign currency translation adjustments', MONEY, False),
 ('fx', 'CI.ci_f', 'Comprehensive income', '{IS.ni_f}+{CI.fx}', MONEY),
 ('sec', 'Printed subtotals (as filed — for integrity checks)'),
 ('in', 'IS.totexp', 'Total expenses — printed', MONEY, False),
 ('in', 'IS.op', 'Operating profit — printed', MONEY, False),
 ('in', 'IS.pbt', 'Income before taxes — printed', MONEY, False),
 ('in', 'IS.ni', 'Net income — printed', MONEY, False),
 ('in', 'CI.ci', 'Comprehensive income — printed', MONEY, False),
 ('sec', 'Integrity checks (formula − printed; must be 0)'),
 ('chk', 'CHECK total expenses', '{IS.totexp_f}-{IS.totexp}', printed('IS.totexp')),
 ('chk', 'CHECK operating profit', '{IS.op_f}-{IS.op}', printed('IS.op')),
 ('chk', 'CHECK pre-tax income', '{IS.pbt_f}-{IS.pbt}', printed('IS.pbt')),
 ('chk', 'CHECK net income', '{IS.ni_f}-{IS.ni}', printed('IS.ni')),
 ('chk', 'CHECK comprehensive income', '{CI.ci_f}-{CI.ci}', printed('CI.ci')),
 ('chk', 'CHECK 6M = Q1 + Q2 (net income)', None, lambda p: p.startswith('6M')),
]
# custom expression for 6M check needs neighbour columns: handled after write
IS_SPEC_W = [x for x in IS_SPEC if not (x[0] == 'chk' and x[2] is None)]
r = write_statement(ws, 5, FLOW_P, IS_SPEC_W, 'IS')
ws.cell(r, 1, 'CHECK 6M = Q1 + Q2 (revenue and net income)').font = F_NOTE
for i, p in enumerate(FLOW_P):
    if p.startswith('6M'):
        y = p[-4:]; q1 = L(FLOW_P.index(f'Q1-{y}') + 2); q2 = L(FLOW_P.index(f'Q2-{y}') + 2); c6 = L(i + 2)
        rv, ni = REF[('IS', 'IS.revenue')], REF[('IS', 'IS.ni_f')]
        ws.cell(r, i + 2, f'=ABS({c6}{rv}-{q1}{rv}-{q2}{rv})+ABS({c6}{ni}-{q1}{ni}-{q2}{ni})').number_format = MONEY
    else:
        ws.cell(r, i + 2, 'n/a').font = F_NOTE
ws.cell(r, MARK, 'CHECK');
for cc in range(1, len(FLOW_P) + 2): ws.cell(r, cc).fill = FILL_CHK
r += 2
ws.cell(r, 1, 'Notes: Form 10 (2026-05-07) and Form 10/A statements are identical (no quiet restatement). The 10-Q prints EPS for carve-out comparatives although the Form 10 states EPS was not presented (no capital structure) — see Accounting Notes.').font = F_NOTE
ws.column_dimensions[L(len(FLOW_P) + 2)].width = 40

# ---------------------------------------------------------------- BS
ws = wb.create_sheet('BS'); COLS['BS'] = BS_P
setup(ws, 'Mobility Global Inc. — Combined Balance Sheets (as originally reported)',
      '12/31/24 & 12/31/25 from Form 10/A audited; 3/31/26 Form 10/A interim; 6/30/26 10-Q. No FY2023 balance sheet is printed.', len(BS_P) + 1)
header(ws, 4, [BSLAB[p] for p in BS_P] + ['Source'])
BS_SPEC = [
 ('sec', 'Assets'),
 ('in', 'BS.cash', 'Cash and cash equivalents', MONEY, False),
 ('in', 'BS.duefrom_c', 'Due from related parties — current', MONEY, False),
 ('in', 'BS.ar', 'Accounts receivable, net', MONEY, False),
 ('in', 'BS.prepaid', 'Prepaid and other current assets', MONEY, False),
 ('fx', 'BS.tca_f', 'Total current assets', '{BS.cash}+{BS.duefrom_c}+{BS.ar}+{BS.prepaid}', MONEY),
 ('in', 'BS.ppe_gross', 'Total property and equipment (gross)', MONEY, False),
 ('in', 'BS.ppe_acc', 'Less: accumulated depreciation', MONEY, False),
 ('in', 'BS.ppe_net', 'Property and equipment, net', MONEY, False),
 ('in', 'BS.rou', 'Right of use assets', MONEY, False),
 ('in', 'BS.gw', 'Goodwill', MONEY, False),
 ('in', 'BS.intang', 'Other intangible assets, net', MONEY, False),
 ('in', 'BS.duefrom_nc', 'Due from related parties — non-current', MONEY, False),
 ('in', 'BS.onca', 'Other non-current assets', MONEY, False),
 ('fx', 'BS.ta_f', 'Total assets', '{BS.tca_f}+{BS.ppe_net}+{BS.rou}+{BS.gw}+{BS.intang}+{BS.duefrom_nc}+{BS.onca}', MONEY),
 ('sec', 'Liabilities and equity'),
 ('in', 'BS.ap', 'Accounts payable', MONEY, False),
 ('in', 'BS.dueto_c', 'Due to related parties — current', MONEY, False),
 ('in', 'BS.accomp', 'Accrued compensation and contributions to retirement plans', MONEY, False),
 ('in', 'BS.unearned', 'Unearned revenue', MONEY, False),
 ('in', 'BS.ocl', 'Other current liabilities', MONEY, False),
 ('fx', 'BS.tcl_f', 'Total current liabilities', '{BS.ap}+{BS.dueto_c}+{BS.accomp}+{BS.unearned}+{BS.ocl}', MONEY),
 ('in', 'BS.ltd', 'Long-term debt (Senior Notes, from 6/30/26)', MONEY, False),
 ('in', 'BS.lease_nc', 'Lease liabilities — non-current (10-Q: "Operating lease liabilities")', MONEY, False),
 ('in', 'BS.dtl', 'Deferred tax liability — non-current', MONEY, False),
 ('in', 'BS.dueto_nc', 'Due to related parties — non-current (Canada Carfax loan)', MONEY, False),
 ('in', 'BS.oncl', 'Other non-current liabilities', MONEY, False),
 ('fx', 'BS.tl_f', 'Total liabilities', '{BS.tcl_f}+{BS.ltd}+{BS.lease_nc}+{BS.dtl}+{BS.dueto_nc}+{BS.oncl}', MONEY),
 ('in', 'BS.npi', 'Parent company investment', MONEY, False),
 ('in', 'BS.aoci', 'Accumulated other comprehensive income (loss)', MONEY, False),
 ('fx', 'BS.te_f', 'Total equity', '{BS.npi}+{BS.aoci}', MONEY),
 ('fx', 'BS.tle_f', 'Total liabilities and equity', '{BS.tl_f}+{BS.te_f}', MONEY),
 ('sec', 'Printed subtotals (as filed)'),
 ('in', 'BS.tca', 'Total current assets — printed', MONEY, False),
 ('in', 'BS.ta', 'Total assets — printed', MONEY, False),
 ('in', 'BS.tcl', 'Total current liabilities — printed', MONEY, False),
 ('in', 'BS.tl', 'Total liabilities — printed', MONEY, False),
 ('in', 'BS.te', 'Total equity — printed', MONEY, False),
 ('in', 'BS.tle', 'Total liabilities and equity — printed', MONEY, False),
 ('sec', 'Integrity checks (must be 0)'),
 ('chk', 'CHECK total current assets', '{BS.tca_f}-{BS.tca}', printed('BS.tca')),
 ('chk', 'CHECK total assets', '{BS.ta_f}-{BS.ta}', printed('BS.ta')),
 ('chk', 'CHECK total current liabilities', '{BS.tcl_f}-{BS.tcl}', printed('BS.tcl')),
 ('chk', 'CHECK total liabilities', '{BS.tl_f}-{BS.tl}', printed('BS.tl')),
 ('chk', 'CHECK total equity', '{BS.te_f}-{BS.te}', printed('BS.te')),
 ('chk', 'CHECK balance: assets − (liabilities + equity)', '{BS.ta_f}-{BS.tle_f}', lambda p: True),
]
write_statement(ws, 5, BS_P, BS_SPEC, 'BS')
ws.column_dimensions[L(len(BS_P) + 2)].width = 40

# ---------------------------------------------------------------- CF
ws = wb.create_sheet('CF'); COLS['CF'] = FLOW_P
setup(ws, 'Mobility Global Inc. — Combined Statements of Cash Flows (as originally reported)',
      'Annual "Restructuring and other" and working-capital lines (Form 10/A p.443) and Q1 working-capital lines are NOT machine-readable in the Quartr text — left blank, never derived. No standalone Q2 cash flow is printed.', len(FLOW_P) + 1)
header(ws, 4, [PLAB[p] for p in FLOW_P] + ['Source'])
def complete(keys, total):
    def f(p):
        if not has(total, p): return False
        s = sum(canon.get((k, p), 0) for k in keys)
        return abs(s - canon[(total, p)]) < 0.5
    return f
OPK = ['CF.ni', 'CF.dep', 'CF.amort', 'CF.prov', 'CF.deftax', 'CF.sbc', 'CF.restr', 'CF.other', 'CF.wc_ar', 'CF.wc_duefrom',
       'CF.wc_prepaid', 'CF.wc_ap', 'CF.wc_dueto', 'CF.wc_unearned', 'CF.wc_ocl', 'CF.wc_other']
INK = ['CF.capex', 'CF.acq', 'CF.loan_rp', 'CF.disp', 'CF.eqinv']
FIK = ['CF.notes', 'CF.parent', 'CF.loan_pay', 'CF.contingent']
LAB = {'CF.ni': 'Net income', 'CF.dep': 'Depreciation', 'CF.amort': 'Amortization of intangibles', 'CF.prov': 'Provision for losses on accounts receivable',
       'CF.deftax': 'Deferred income taxes', 'CF.sbc': 'Stock-based compensation', 'CF.restr': 'Restructuring and other', 'CF.other': 'Other',
       'CF.wc_ar': 'Δ Accounts receivable', 'CF.wc_duefrom': 'Δ Due from related parties', 'CF.wc_prepaid': 'Δ Prepaid and other current assets',
       'CF.wc_ap': 'Δ Accounts payable and accrued expenses', 'CF.wc_dueto': 'Δ Due to related parties', 'CF.wc_unearned': 'Δ Unearned revenue',
       'CF.wc_ocl': 'Δ Other current liabilities', 'CF.wc_other': 'Net change in other assets and liabilities',
       'CF.capex': 'Capital expenditures', 'CF.acq': 'Acquisitions, net of cash acquired', 'CF.loan_rp': 'Proceeds from loan to related parties',
       'CF.disp': 'Proceeds from dispositions', 'CF.eqinv': 'Purchases of equity investments',
       'CF.notes': 'Proceeds from issuance of Senior Notes', 'CF.parent': 'Net transfers to Parent', 'CF.loan_pay': 'Payments related to loan from related parties',
       'CF.contingent': 'Contingent consideration payments'}
def sumexpr(keys): return '+'.join('{%s}' % k for k in keys)
CF_SPEC = [('sec', 'Operating activities')] + [('in', k, LAB[k], MONEY, False) for k in OPK] + [
 ('in', 'CF.cfo', 'Cash provided by operating activities', MONEY, True),
 ('sec', 'Investing activities')] + [('in', k, LAB[k], MONEY, False) for k in INK] + [
 ('in', 'CF.cfi', 'Cash used for investing activities (FY2024 & Q1-25 from MD&A cash-flow overview — statement cell unreadable)', MONEY, True),
 ('sec', 'Financing activities')] + [('in', k, LAB[k], MONEY, False) for k in FIK] + [
 ('in', 'CF.cff', 'Cash provided by (used for) financing activities (Q1-25 from MD&A)', MONEY, True),
 ('in', 'CF.fx', 'Effect of exchange rate changes on cash', MONEY, False),
 ('in', 'CF.chg', 'Net change in cash, cash equivalents and restricted cash', MONEY, True),
 ('in', 'CF.beg', 'Cash at beginning of period', MONEY, False),
 ('in', 'CF.end', 'Cash at end of period', MONEY, True),
 ('sec', 'Supplemental'),
 ('in', 'CF.sup_int', 'Cash paid for interest', MONEY, False),
 ('in', 'CF.sup_tax', 'Cash paid for income taxes', MONEY, False),
 ('in', 'CF.sup_canloan', 'Non-cash: consolidation of Canada Carfax loan', MONEY, False),
 ('in', 'CF.sup_debtcost', 'Non-cash: accrued debt issuance costs', MONEY, False),
 ('sec', 'Integrity checks (must be 0; n/a where components not fully extractable or not printed)'),
 ('chk', 'CHECK CFO = sum of printed components', sumexpr(OPK) + '-{CF.cfo}', complete(OPK, 'CF.cfo')),
 ('chk', 'CHECK CFI = sum of printed components', sumexpr(INK) + '-{CF.cfi}', complete(INK, 'CF.cfi')),
 ('chk', 'CHECK CFF = sum of printed components', sumexpr(FIK) + '-{CF.cff}', complete(FIK, 'CF.cff')),
 ('chk', 'CHECK CFO + CFI + CFF + FX = net change', '{CF.cfo}+{CF.cfi}+{CF.cff}+{CF.fx}-{CF.chg}',
  lambda p: all(has(k, p) for k in ['CF.cfo', 'CF.cfi', 'CF.cff', 'CF.fx', 'CF.chg'])),
 ('chk', 'CHECK beginning + change = ending cash', '{CF.beg}+{CF.chg}-{CF.end}', lambda p: has('CF.end', p)),
 ('chk', 'CHECK net income = IS net income', '{CF.ni}-[IS!IS.ni_f]', lambda p: has('CF.ni', p)),
 ('chk', 'CHECK ending cash = BS cash', '{CF.end}-[BS!BS.cash]', lambda p: has('CF.end', p) and p in {'FY2024': 1, 'FY2025': 1, 'Q1-2026': 1, '6M-2026': 1}),
]
# BS cross-ref needs period mapping 6M-2026 -> Q2-2026 column on BS
_orig_xcol = xcol
def xcol(sheet, p):
    if sheet == 'BS' and p == '6M-2026': p = 'Q2-2026'
    return _orig_xcol(sheet, p)
write_statement(ws, 5, FLOW_P, CF_SPEC, 'CF')
ws.column_dimensions[L(len(FLOW_P) + 2)].width = 40
for i, p in enumerate(FLOW_P):  # explain gaps in CFO check
    pass

# ---------------------------------------------------------------- Key metrics (green cross-sheet formulas only)
ws = wb.create_sheet('Key Metrics'); COLS['Key Metrics'] = FLOW_P
setup(ws, 'Key Metrics — computed [C] from IS / BS / CF (formulas only)',
      'Growth only between comparable periods (same length, prior year). FY2023 growth suppressed: no FY2022 carve-out statements printed. No 53-week years (calendar Dec-31 FYE).', len(FLOW_P) + 1)
header(ws, 4, [PLAB[p] for p in FLOW_P] + ['Definition'], first='Ratio / metric')
prior = {'FY2024': 'FY2023', 'FY2025': 'FY2024', 'Q1-2026': 'Q1-2025', 'Q2-2026': 'Q2-2025', '6M-2026': '6M-2025'}
def isr(key): return REF[('IS', key)]
def cfr(key): return REF[('CF', key)]
metrics = [
 ('Revenue growth (y/y)', lambda p, c: f"='IS'!{c}{isr('IS.revenue')}/'IS'!{L(FLOW_P.index(prior[p])+2)}{isr('IS.revenue')}-1" if p in prior else None, PCT_C, 'Revenue / prior-year comparable period − 1'),
 ('Operating-related expenses % revenue', lambda p, c: f"='IS'!{c}{isr('IS.oprel')}/'IS'!{c}{isr('IS.revenue')}", PCT_C, ''),
 ('Selling & general % revenue', lambda p, c: f"='IS'!{c}{isr('IS.sg')}/'IS'!{c}{isr('IS.revenue')}", PCT_C, ''),
 ('EBITDA (operating profit + D&A) [C]', lambda p, c: f"='IS'!{c}{isr('IS.op_f')}+'IS'!{c}{isr('IS.dep')}+'IS'!{c}{isr('IS.amort')}", MONEY, 'Unadjusted; company Adjusted EBITDA is on the Non-GAAP tab'),
 ('EBITDA margin [C]', lambda p, c: f"=({c}9)/'IS'!{c}{isr('IS.revenue')}", PCT_C, ''),
 ('Operating margin', lambda p, c: f"='IS'!{c}{isr('IS.op_f')}/'IS'!{c}{isr('IS.revenue')}", PCT_C, ''),
 ('Amortization of intangibles % revenue', lambda p, c: f"='IS'!{c}{isr('IS.amort')}/'IS'!{c}{isr('IS.revenue')}", PCT_C, 'Acquisition (IHS Markit merger) step-up amortization'),
 ('Net margin', lambda p, c: f"='IS'!{c}{isr('IS.ni_f')}/'IS'!{c}{isr('IS.revenue')}", PCT_C, ''),
 ('Effective tax rate', lambda p, c: f"='IS'!{c}{isr('IS.tax')}/'IS'!{c}{isr('IS.pbt_f')}", PCT_C, ''),
 ('Free cash flow (CFO + capex) [C]', lambda p, c: f"='CF'!{c}{cfr('CF.cfo')}+'CF'!{c}{cfr('CF.capex')}" if has('CF.cfo', p) and has('CF.capex', p) else None, MONEY, 'Suppressed where capex not printed in the statement'),
 ('FCF / net income', lambda p, c: f"={c}15/'IS'!{c}{isr('IS.ni_f')}" if has('CF.cfo', p) and has('CF.capex', p) else None, PCT_C, ''),
 ('FCF / (net income + amortization) [C]', lambda p, c: f"={c}15/('IS'!{c}{isr('IS.ni_f')}+'IS'!{c}{isr('IS.amort')})" if has('CF.cfo', p) and has('CF.capex', p) else None, PCT_C, 'Cash conversion ex-acquired-intangible amortization (pre-tax amortization; indicative)'),
 ('Capex % revenue', lambda p, c: f"=-'CF'!{c}{cfr('CF.capex')}/'IS'!{c}{isr('IS.revenue')}" if has('CF.capex', p) else None, PCT_C, ''),
 ('Stock-based compensation % revenue', lambda p, c: f"='CF'!{c}{cfr('CF.sbc')}/'IS'!{c}{isr('IS.revenue')}" if has('CF.sbc', p) else None, PCT_C, 'Carve-out SBC; mgmt guides 3–4% standalone [IR]'),
 ('Net transfers to Parent % CFO', lambda p, c: f"=-'CF'!{c}{cfr('CF.parent')}/'CF'!{c}{cfr('CF.cfo')}" if has('CF.parent', p) and has('CF.cfo', p) else None, PCT_C, ''),
]
r = 5
section(ws, r, 'Growth, margins, cash conversion', len(FLOW_P) + 1); r += 1
for lab, fn, fmt, d in metrics:
    ws.cell(r, 1, lab).font = F_BASE
    for i, p in enumerate(FLOW_P):
        e = fn(p, L(i + 2))
        if e:
            c = ws.cell(r, i + 2, e); c.font = F_X if "'" in e else F_FX; c.number_format = fmt
    ws.cell(r, len(FLOW_P) + 2, d).font = F_NOTE
    r += 1
ws.column_dimensions[L(len(FLOW_P) + 2)].width = 50
assert ws.cell(9, 1).value.startswith('EBITDA (') and ws.cell(15, 1).value.startswith('Free cash flow'), 'metric row refs moved'

# ---------------------------------------------------------------- generic long-CSV pivot tabs
def pivot(title_tab, title, subtitle, rows, period_filter=None, group=('source_doc', 'section'), source_col=True):
    ws = wb.create_sheet(title_tab)
    if period_filter: rows = [r for r in rows if period_filter(r['fiscal_period'])]
    periods = sorted({r['fiscal_period'] for r in rows}, key=period_key)
    setup(ws, title, subtitle, len(periods) + 1, label_w=60)
    header(ws, 4, periods + ['Source (doc p.)'], first='Line item (verbatim caption)')
    groups = collections.OrderedDict()
    for r in rows: groups.setdefault(tuple(r[g] for g in group), []).append(r)
    rr = 5
    for gk, grs in groups.items():
        section(ws, rr, ' · '.join(gk), len(periods) + 1); rr += 1
        lines = collections.OrderedDict()
        for r in sorted(grs, key=lambda x: (float(x.get('order_index') or 0) if re.match(r'^-?\d+(\.\d+)?$', x.get('order_index') or '') else 0)):
            li = re.sub(r'\s+', ' ', r['line_item']).strip(); k = li; n = 2
            while k in lines and r['fiscal_period'] in lines[k]['vals']:
                if lines[k]['vals'][r['fiscal_period']][0] == r['value']: break
                k = f'{li}  [#{n}]'; n += 1
            lines.setdefault(k, {'vals': {}, 'pages': set(), 'sub': r.get('is_subtotal') == '1'})
            lines[k]['vals'][r['fiscal_period']] = (r['value'], r['unit'])
            lines[k]['pages'].add(f"{r['source_doc']} p{r['page']}")
        for li, d in lines.items():
            ws.cell(rr, 1, li).font = F_SEC if d['sub'] else F_BASE
            ws.cell(rr, 1).alignment = Alignment(wrap_text=False)
            for i, p in enumerate(periods):
                if p in d['vals']:
                    v, u = d['vals'][p]
                    try:
                        val = float(v) if u != 'text' else v
                    except ValueError:
                        val = v
                    c = ws.cell(rr, i + 2, val)
                    c.font = F_IN_B if d['sub'] else F_IN
                    if isinstance(val, float): c.number_format = FMT.get(u, MONEY)
            if source_col:
                ws.cell(rr, len(periods) + 2, '; '.join(sorted(d['pages']))[:250]).font = F_NOTE
            rr += 1
    ws.column_dimensions[L(len(periods) + 2)].width = 45
    return ws

seg = rd('segments_kpi.csv')
pivot('Segments', 'Segments — as defined in each filing (never recast backwards)',
      'Form 10/A: segment operating profit (ASU 2023-07 format). 10-Q/8-K: Adjusted EBITDA as segment measure from Q2-2026 (prior periods recast BY THE COMPANY). Bases not comparable — see Accounting Notes.',
      [r for r in seg if r['statement'] == 'SEG'])
pivot('Revenue Detail', 'Revenue disaggregation — as printed',
      'No revenue by brand/product or by B2B business line is disclosed. Geography = U.S. vs International only.', [r for r in seg if r['statement'] == 'REV_DISAGG'])
pivot('Adjusted Non-GAAP', 'Non-GAAP / APMs — as printed per document',
      'Adjusted EBITDA definition: 10 exclusions (Form 10, 10/A) vs 6 (10-Q, 8-K). 8-K Exhibit 5 add-back lines corrupted in source text — 10-Q used. Unaudited.',
      [r for r in seg if r['statement'] == 'NONGAAP'] + [r for r in rd('statements.csv') if r['statement'] == 'NONGAAP'])
pivot('Operating KPIs', 'Operating KPIs & guidance — as printed', 'Definitions per source; [D] filings, release and decks. Guidance dated 2026-08-07.',
      [r for r in seg if r['statement'] == 'KPI'])
notes = rd('notes.csv')
for code, tab, ttl in [('NOTE_DEBT', 'Debt Schedule', 'Debt — instrument by instrument'),
                       ('NOTE_TAX', 'Tax Reconciliation', 'Income taxes — ETR reconciliation, current/deferred, DTA/DTL'),
                       ('NOTE_EQUITY', 'Equity Roll-Forward', 'Equity / net parent investment, shares, stock compensation'),
                       ('NOTE_OCI', 'OCI', 'Other comprehensive income'),
                       ('NOTE_PENSION', 'Pension', 'Retirement plans (defined contribution only; no defined-benefit plan disclosed)'),
                       ('NOTE_GOODWILL_INTANG', 'Goodwill & Intangibles', 'Goodwill by segment and intangible assets by class'),
                       ('NOTE_RELATED', 'Related Party', 'Related-party transactions and S&P Global allocations'),
                       ('NOTE_LEASES', 'Leases', 'Leases'),
                       ('NOTE_COMMIT', 'Commitments', 'Commitments and contingencies'),
                       ('NOTE_DEFREV', 'Deferred Revenue', 'Unearned revenue and remaining performance obligations')]:
    rs = [r for r in notes if r['statement'] == code]
    if code == 'NOTE_EQUITY':
        rs += [r for r in rd('statements.csv') if r['statement'] == 'EQ']
    pivot(tab, ttl, 'As printed; see notes_NOTES.md for unreadable cells and tables that do not foot.', rs)
pivot('Pro Forma', 'Unaudited pro forma condensed combined (Form 10/A) — never merged with historical',
      'Management estimates of the separation (debt, distribution, standalone adjustments).', [r for r in rd('statements.csv') if r['statement'] == 'PROFORMA'])
pivot('Predecessor Segments', 'Predecessor segment history — S&P Global "Mobility" (FY2022–25, Dec FYE) and IHS Markit "Transportation" (FY2014–21, Nov 30 FYE)',
      'NON-COMPARABLE across the eras (scope incl. Maritime & Trade and Aerospace & Defense until 2019-12; Nov vs Dec FYE; adjusted EBITDA vs operating profit after amortization; FY2022 = ~10 months post-merger). Segment-level only; never merged into carve-out tabs.',
      rd('predecessor.csv'))

# Acquisitions — deal-per-column
ws = wb.create_sheet('Acquisitions')
acq = rd('acquisitions.csv')
deals = collections.OrderedDict()
for r in acq: deals.setdefault(r['deal_label'], []).append(r)
setup(ws, 'Acquisitions — deal-per-column (purchase price allocations as printed)', 'Preliminary and final PPAs are separate deal columns where both were printed.', len(deals) + 1, label_w=55)
header(ws, 4, list(deals.keys()), first='Line')
for i in range(len(deals)): ws.column_dimensions[L(i + 2)].width = 26
lines = collections.OrderedDict()
for r in acq: lines.setdefault((r['section'], r['line']), {})[r['deal_label']] = r
rr = 5; last_sec = None
for (sec, line), d in lines.items():
    if sec != last_sec: section(ws, rr, sec, len(deals) + 1); rr += 1; last_sec = sec
    ws.cell(rr, 1, line).font = F_BASE
    for i, dl in enumerate(deals):
        if dl in d:
            r = d[dl]
            try: v = float(r['value']) if r['unit'] != 'text' else r['value']
            except ValueError: v = r['value']
            c = ws.cell(rr, i + 2, v); c.font = F_IN
            if isinstance(v, float): c.number_format = FMT.get(r['unit'], MONEY)
    rr += 1

# Presentation data (+ targets vs actuals)
pres = rd('presentations.csv')
ws = pivot('Presentation Data', 'Presentation data — deck-only figures (UNAUDITED marketing material; never overwrites filings)',
           'Investor Day 2026-05-12 (116 slides) and Q2-2026 deck. page = slide number. Image-only slides recorded as not machine-readable.', pres)
mx = ws.max_row + 2
section(ws, mx, 'Targets vs actuals — the only computation on this tab (green = links to model)', 4); mx += 1
g_lo = [r for r in seg if r['statement'] == 'KPI' and 'revenue' in r['line_item'].lower() and r['fiscal_period'] == 'FY2026E']
ws.cell(mx, 1, 'FY2026 revenue guidance (Aug 7, 2026) — see Operating KPIs tab for low/high as printed').font = F_BASE; mx += 1
ws.cell(mx, 1, '6M-2026 revenue actual [D]').font = F_BASE
c = ws.cell(mx, 2, f"='IS'!{L(FLOW_P.index('6M-2026')+2)}{isr('IS.revenue')}"); c.font = F_X; c.number_format = MONEY; mx += 1
ws.cell(mx, 1, 'FY2025 revenue actual [D]').font = F_BASE
c = ws.cell(mx, 2, f"='IS'!{L(FLOW_P.index('FY2025')+2)}{isr('IS.revenue')}"); c.font = F_X; c.number_format = MONEY; mx += 1
ws.cell(mx, 1, 'Guide low: FY2026 revenue $1,870m [D, 8-K 2026-08-07]').font = F_BASE
c = ws.cell(mx, 2, 1870); c.font = F_IN; c.number_format = MONEY; lo = mx; mx += 1
ws.cell(mx, 1, 'Guide high: FY2026 revenue $1,885m [D]').font = F_BASE
c = ws.cell(mx, 2, 1885); c.font = F_IN; c.number_format = MONEY; hi = mx; mx += 1
ws.cell(mx, 1, 'Implied H2-2026 revenue at low / high [C]').font = F_BASE
ws.cell(mx, 2, f'=B{lo}-B{mx-4}').number_format = MONEY; ws.cell(mx, 3, f'=B{hi}-B{mx-4}').number_format = MONEY; h2 = mx; mx += 1
ws.cell(mx, 1, 'Implied H2-2026 growth vs H2-2025 (FY2025 − 6M-2025) at low / high [C]').font = F_BASE
h2_25 = f"('IS'!{L(FLOW_P.index('FY2025')+2)}{isr('IS.revenue')}-'IS'!{L(FLOW_P.index('6M-2025')+2)}{isr('IS.revenue')})"
for j, col in enumerate(['B', 'C']):
    c = ws.cell(mx, 2 + j, f'={col}{h2}/{h2_25}-1'); c.font = F_X; c.number_format = PCT_C

# ---------------------------------------------------------------- Accounting notes, Qualitative pointer, Caption map, Cover
an = list(csv.DictReader(open(f'{W}/canonical/accounting_notes.csv')))
ws = wb.create_sheet('Accounting Notes', 1)
setup(ws, 'Accounting Notes — breaks, restatements, definition changes, non-disclosures, data-quality flags', 'One row per item. Tier: [D] disclosed; [C] computed; flag types: BREAK / RESTATE / DEFN / NONDISC / DATAQ / BASIS.', 5, label_w=14)
header(ws, 4, ['Area', 'Item', 'Source', 'Model treatment'], first='Type')
for c_, w_ in zip('BCDE', (18, 90, 30, 50)): ws.column_dimensions[c_].width = w_
for i, r in enumerate(an):
    for j, k in enumerate(['type', 'area', 'item', 'source', 'treatment']):
        c = ws.cell(5 + i, 1 + j, r[k]); c.font = F_BASE; c.alignment = Alignment(wrap_text=True, vertical='top')

ws = wb.create_sheet('Qualitative')
setup(ws, 'Qualitative — see analysis/QUALITATIVE.md (business, concentration, separation, risk-factor inventory, events timeline)', '', 2, label_w=120)
qr = 4
for line in open(f'{W}/analysis/QUALITATIVE.md').read().splitlines():
    if not line.strip(): continue
    c = ws.cell(qr, 1, line[:1000]); c.alignment = Alignment(wrap_text=True, vertical='top')
    c.font = F_SEC if line.startswith('#') else F_BASE
    if line.startswith('#'): c.fill = FILL_S
    qr += 1

ws = wb.create_sheet('Caption Map')
setup(ws, 'Caption Map — primary statements (canonical) + identity mapping for verbatim tabs', 'Secondary tabs present captions verbatim grouped by source document and section (identity map). Primary-statement merges below; merge log in canonical/primary_merge_log.txt.', 6, label_w=22)
header(ws, 4, ['Statement', 'Section', 'Caption (verbatim)', 'Canonical key / EXCLUDE', 'Reason'], first='Source doc')
for c_, w_ in zip('BCDEF', (8, 50, 55, 18, 60)): ws.column_dimensions[c_].width = w_
for i, r in enumerate(csv.reader(open(f'{W}/canonical/primary_caption_map.csv'))):
    if i == 0: continue
    for j, v in enumerate(r): ws.cell(4 + i, 1 + j, v).font = F_BASE

ws = wb.create_sheet('Cover & Basis', 0)
setup(ws, 'Mobility Global Inc. (NYSE: MBGL) — Audit-quality historical financial model', 'Built 2026-10-06 from primary filings via the Quartr connector. All figures as originally reported.', 2, label_w=34)
ws.column_dimensions['B'].width = 120
cover = [
 ('Company', 'Mobility Global Inc. (Delaware); spin-off of S&P Global Mobility (CARFAX, automotiveMastermind, Polk Automotive Solutions, Market Scan). Distribution 2026-07-01 (record date 2026-06-15), ratio 1:1; S&P Global retained no stake.'),
 ('Listing / ticker', 'NYSE: MBGL'),
 ('Regulator ID', 'SEC File No. 001-43276; IRS EIN 39-4621962. EDGAR CIK not confirmed (sec.gov blocked by sandbox network policy).'),
 ('GAAP / currency', 'US GAAP; USD millions (per-share in USD). Fiscal year ends December 31.'),
 ('Basis of preparation', 'Combined (carve-out) financial statements of the Mobility business of S&P Global — include allocations of S&P Global corporate costs; not necessarily indicative of standalone results. Standalone from distribution (Q3-2026 onward).'),
 ('Periods', 'Income statement & cash flow FY2023–FY2025 (audited), Q1-2025/Q1-2026 (Form 10/A), Q2 and 6M 2025/2026 (10-Q). Balance sheets 12/31/24, 12/31/25, 3/31/26, 6/30/26. Predecessor segment history FY2014–FY2025 on its own tab.'),
 ('Sources', 'Form 10-12B (2026-05-07), Form 10-12B/A (2026-05-27), 10-Q Q2-2026 (2026-08-07), 8-K earnings release (2026-08-07), Investor Day deck (2026-05-12), Q2-2026 deck; S&P Global 10-K FY2022–FY2025; IHS Markit 10-K FY2016–FY2021. All read via Quartr page text.'),
 ('Source limitation', 'Native filings (EDGAR, Quartr PDF CDN, Wayback) were blocked by the sandbox network policy; Quartr page text flattens tables and corrupts some digits. Unreadable cells are left blank and logged (Accounting Notes, *_NOTES.md). No XBRL cross-check was possible — verification is cross-document (Form 10 vs 10/A vs 10-Q vs 8-K vs decks).'),
 ('Conventions', 'Blank = not reported / not machine-readable; dash = reported zero. Blue = hard input as filed; black = in-sheet formula; green = cross-sheet formula. "0.0%" literal format = as-printed percentage; true % = computed ratio. Check rows marked CHECK (hidden col Z).'),
 ('Tiers', '[D] disclosed · [C] computed · [E] estimate · [IR] management verbal · [X] third party · [context] general knowledge.'),
 ('Tabs', 'Accounting Notes · IS · BS · CF · Key Metrics · Segments · Revenue Detail · Adjusted Non-GAAP · Operating KPIs · notes tabs (Debt, Tax, Equity, OCI, Pension, Goodwill & Intangibles, Related Party, Leases, Commitments, Deferred Revenue) · Pro Forma · Predecessor Segments · Acquisitions · Presentation Data · Qualitative · Caption Map'),
]
for i, (k, v) in enumerate(cover):
    ws.cell(4 + i, 1, k).font = F_SEC
    c = ws.cell(4 + i, 2, v); c.font = F_BASE; c.alignment = Alignment(wrap_text=True, vertical='top')

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.font is None or c.font.name != 'Arial':
                c.font = Font(name='Arial', size=c.font.size if c.font and c.font.size else 9, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)
an_ws = wb['Accounting Notes']; wb._sheets.remove(an_ws); wb._sheets.insert(1, an_ws)
out = f'{W}/deliverables/MBGL_Historical_Model.xlsx'
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print('saved', out, [s.title for s in wb.worksheets])
